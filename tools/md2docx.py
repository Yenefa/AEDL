# -*- coding: utf-8 -*-
"""md2docx.py — 把本项目 markdown 子集转成规范 Word 报告。

支持语法：# 标题(1-4 级)、段落、**加粗**、| 表格 |、- 无序列表、1. 有序列表、
> 引用、``` 代码块、--- 分隔线（忽略）。

用法:
    python tools/md2docx.py docs/report-draft.md "docs/报告.docx"
退出码: 0 成功; 非 0 失败（真实退出码显式传播，禁止被后续命令掩盖）。
"""
import re
import sys

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    sys.stderr.write("FATAL: python-docx 未安装 (pip install python-docx)\n")
    sys.exit(2)


def set_cn_font(run, name="微软雅黑", size=None, bold=None, mono=False):
    if mono:
        name = "Consolas"
    run.font.name = name
    r = run._element.rPr
    rFonts = r.rFonts if r.rFonts is not None else r.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "微软雅黑" if not mono else name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_runs(par, text, base_size=10.5, mono=False):
    """把 **bold** 内联标记拆成多个 run。"""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            set_cn_font(r, size=base_size, mono=mono)
        rb = par.add_run(m.group(1))
        set_cn_font(rb, size=base_size, bold=True, mono=mono)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        set_cn_font(r, size=base_size, mono=mono)


def add_table(doc, rows):
    """rows: list[list[str]]，第一行为表头。"""
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            txt = row[j] if j < len(row) else ""
            add_runs(p, txt, base_size=9)
            if i == 0:  # 表头加粗+灰底
                for r in p.runs:
                    r.font.bold = True
                from docx.oxml import OxmlElement
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "E7E6E6")
                tcPr.append(shd)


def convert(md_path, out_path):
    with open(md_path, "r", encoding="utf-8-sig") as f:
        lines = f.read().splitlines()

    doc = Document()
    # 正文默认字体
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "微软雅黑")

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            for bl in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(bl if bl.strip() else " ")
                set_cn_font(r, size=9, mono=True)
            continue

        # 表格块
        if stripped.startswith("|") and stripped.endswith("|"):
            tbl_rows = []
            while i < n and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    tbl_rows.append(cells)
                i += 1
            if tbl_rows:
                add_table(doc, tbl_rows)
                doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1)) - 1  # H1→Heading0 用 Title 处理
            text = m.group(2).strip()
            if level == 0:
                p = doc.add_heading("", level=0)
                add_runs(p, text, base_size=20)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                hl = min(level, 3)
                p = doc.add_paragraph()
                p.style = doc.styles[f"Heading {hl}"]
                add_runs(p, text, base_size={1: 16, 2: 14, 3: 12}[hl])
            i += 1
            continue

        # 分隔线：跳过
        if re.fullmatch(r"-{3,}", stripped):
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            add_runs(p, stripped.lstrip("> ").strip(), base_size=10)
            for r in p.runs:
                r.font.italic = True
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            i += 1
            continue

        # 无序列表
        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^[-*]\s+", "", stripped), base_size=10.5)
            i += 1
            continue

        # 有序列表
        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", stripped), base_size=10.5)
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, stripped, base_size=10.5)
        i += 1

    doc.save(out_path)


def main(argv):
    if len(argv) != 3:
        sys.stderr.write("usage: python tools/md2docx.py <input.md> <output.docx>\n")
        return 2
    convert(argv[1], argv[2])
    print(f"OK: {argv[2]}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
